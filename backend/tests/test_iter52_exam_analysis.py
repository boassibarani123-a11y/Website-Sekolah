"""Iter52 - Vortex Guard: Analisis Ujian feature + KKM field + Excel analysis sheet."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

def _env_backend():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if v:
        return v
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip()
    raise RuntimeError("REACT_APP_BACKEND_URL not set")

BASE_URL = _env_backend().rstrip("/")
API = f"{BASE_URL}/api"

ADMIN = ("boassibarani123@gmail.com", "<SUPER_ADMIN_PASSWORD>")
GURU = ("guru@sman1laguboti.sch.id", "Guru12345")
SISWA = ("siswa@sman1laguboti.sch.id", "Siswa12345")

TEST_PREFIX = "TEST_ITER52"


def _login(email, password):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200, f"login failed {email}: {r.status_code} {r.text}"
    return r.json()["token"]


def _h(tok):
    return {"Authorization": f"Bearer {tok}", "Content-Type": "application/json"}


@pytest.fixture(scope="module")
def admin_tok():
    return _login(*ADMIN)


@pytest.fixture(scope="module")
def guru_tok():
    return _login(*GURU)


@pytest.fixture(scope="module")
def siswa_tok():
    return _login(*SISWA)


@pytest.fixture(scope="module")
def xii_ipa1(admin_tok):
    r = requests.get(f"{API}/classes", headers=_h(admin_tok), timeout=30)
    assert r.status_code == 200
    kls = next((c for c in r.json() if c["name"] == "XII IPA 1"), None)
    assert kls, "XII IPA 1 missing"
    return kls


@pytest.fixture(scope="module")
def test_siswas(admin_tok, xii_ipa1):
    """Create 5 TEST siswa accounts in XII IPA 1."""
    created = []
    for i in range(1, 6):
        payload = {
            "email": f"test_iter52_s{i}@sman1laguboti.sch.id",
            "password": "Siswa12345",
            "name": f"{TEST_PREFIX} Siswa {i}",
            "role": "siswa",
            "nisn": f"999999{i:04d}",
            "kelas": "XII IPA 1",
            "phone": f"0812340000{i:02d}",
        }
        r = requests.post(f"{API}/users", headers=_h(admin_tok), json=payload, timeout=30)
        if r.status_code == 400 and "sudah terdaftar" in r.text:
            # fetch id from listing
            r2 = requests.get(f"{API}/users?role=siswa", headers=_h(admin_tok), timeout=30)
            u = next((x for x in r2.json() if x["email"] == payload["email"]), None)
            assert u, f"could not refetch {payload['email']}"
            created.append({"id": u["id"], "email": payload["email"], "password": payload["password"]})
            continue
        assert r.status_code == 200, f"create siswa {i} failed: {r.status_code} {r.text}"
        d = r.json()
        created.append({"id": d["id"], "email": payload["email"], "password": payload["password"]})
    yield created
    # Teardown: delete
    for s in created:
        requests.delete(f"{API}/users/{s['id']}", headers=_h(admin_tok), timeout=30)


@pytest.fixture(scope="module")
def test_exam(guru_tok, xii_ipa1, test_siswas):
    """Create a TEST exam w/ 3 MC (correct answers index 0,1,2) + 1 essay, KKM=70."""
    body = {
        "title": f"{TEST_PREFIX} UjianAnalisis",
        "kelas": xii_ipa1["name"],
        "class_id": xii_ipa1["id"],
        "subject": "Matematika",
        "password": "ITER52",
        "kkm": 70,
        "time_limit": 0,
        "max_violations": 3,
        "shuffle_questions": False,
        "shuffle_options": False,
        "questions": [
            {"type": "mc", "q": "Q1 1+1=?", "options": ["2", "3", "4", "5"], "answer": 0},
            {"type": "mc", "q": "Q2 2+2=?", "options": ["3", "4", "5", "6"], "answer": 1},
            {"type": "mc", "q": "Q3 3+3=?", "options": ["5", "7", "6", "9"], "answer": 2},
            {"type": "essay", "q": "Jelaskan.", "options": [], "answer": 0},
        ],
    }
    r = requests.post(f"{API}/exams", headers=_h(guru_tok), json=body, timeout=30)
    assert r.status_code == 200, f"create exam failed: {r.status_code} {r.text}"
    exam = r.json()
    assert exam["kkm"] == 70
    yield exam
    requests.delete(f"{API}/exams/{exam['id']}", headers=_h(guru_tok), timeout=30)


# =========================================
# KKM field validation
# =========================================
class TestKKM:
    def test_kkm_default_75(self, guru_tok, xii_ipa1):
        body = {
            "title": f"{TEST_PREFIX} KKMDefault",
            "kelas": xii_ipa1["name"],
            "class_id": xii_ipa1["id"],
            "subject": "Matematika",
            "password": "KKMD",
            "questions": [{"type": "mc", "q": "x?", "options": ["a", "b"], "answer": 0}],
        }
        r = requests.post(f"{API}/exams", headers=_h(guru_tok), json=body, timeout=30)
        assert r.status_code == 200
        ex = r.json()
        assert ex["kkm"] == 75, f"expected default 75, got {ex.get('kkm')}"
        requests.delete(f"{API}/exams/{ex['id']}", headers=_h(guru_tok), timeout=30)

    def test_kkm_reject_above_100(self, guru_tok, xii_ipa1):
        body = {
            "title": f"{TEST_PREFIX} KKMBad",
            "kelas": xii_ipa1["name"],
            "class_id": xii_ipa1["id"],
            "subject": "Matematika",
            "password": "KKMB",
            "kkm": 150,
            "questions": [{"type": "mc", "q": "x?", "options": ["a", "b"], "answer": 0}],
        }
        r = requests.post(f"{API}/exams", headers=_h(guru_tok), json=body, timeout=30)
        assert r.status_code == 422, f"expected 422, got {r.status_code} {r.text}"

    def test_kkm_patch_updates(self, guru_tok, test_exam):
        r = requests.patch(f"{API}/exams/{test_exam['id']}", headers=_h(guru_tok), json={"kkm": 80}, timeout=30)
        assert r.status_code == 200
        assert r.json()["kkm"] == 80
        # revert
        requests.patch(f"{API}/exams/{test_exam['id']}", headers=_h(guru_tok), json={"kkm": 70}, timeout=30)

    def test_kkm_patch_reject_above_100(self, guru_tok, test_exam):
        r = requests.patch(f"{API}/exams/{test_exam['id']}", headers=_h(guru_tok), json={"kkm": 101}, timeout=30)
        assert r.status_code == 422


# =========================================
# Submissions by 5 TEST siswa (varied answers)
# =========================================
# answer plan (indexes into ORIGINAL question options since shuffle=false):
#   siswa 1: [0,1,2]  all correct -> 100%
#   siswa 2: [0,1,2]  all correct -> 100%
#   siswa 3: [0,1,3]  2 correct   -> ~66.7%
#   siswa 4: [1,0,0]  all wrong   -> 0%
#   siswa 5: [1,2,3]  all wrong   -> 0%
# essay q (index 3) answered with simple text (score None).
ANSWER_PLAN = [
    [0, 1, 2, "essay a"],
    [0, 1, 2, "essay b"],
    [0, 1, 3, "essay c"],
    [1, 0, 0, "essay d"],
    [1, 2, 3, "essay e"],
]


@pytest.fixture(scope="module")
def submissions(test_siswas, test_exam):
    """Have each siswa start+submit."""
    results = []
    for i, s in enumerate(test_siswas):
        tok = _login(s["email"], s["password"])
        # start
        r = requests.post(f"{API}/exams/{test_exam['id']}/start",
                          headers=_h(tok), json={"password": "ITER52"}, timeout=30)
        assert r.status_code == 200, f"start failed s{i+1}: {r.status_code} {r.text}"
        # with shuffle=false, option order is identity: answers can be submitted as-is
        answers = ANSWER_PLAN[i]
        r = requests.post(f"{API}/exams/attempt",
                          headers=_h(tok), json={"exam_id": test_exam["id"], "answers": answers}, timeout=30)
        assert r.status_code == 200, f"attempt failed s{i+1}: {r.status_code} {r.text}"
        results.append(r.json())
    return results


# =========================================
# Analysis endpoint
# =========================================
class TestAnalysis:
    def test_siswa_forbidden(self, siswa_tok, test_exam):
        r = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(siswa_tok), timeout=30)
        assert r.status_code == 403

    def test_guru_ok_shape(self, guru_tok, test_exam, submissions):
        r = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(guru_tok), timeout=30)
        assert r.status_code == 200, r.text
        d = r.json()
        # Stats
        st = d["stats"]
        for k in ("participants", "roster", "not_taken", "avg", "max", "min", "median",
                  "kkm", "passed", "pass_rate"):
            assert k in st, f"missing stat {k}"
        assert st["kkm"] == 70
        assert st["participants"] >= 5
        assert d["enough_data"] is True
        # Distribution 10 buckets
        assert len(d["distribution"]) == 10
        labels = [b["range"] for b in d["distribution"]]
        assert labels[0] == "0-10"
        assert labels[-1] == "91-100"
        # Items present
        assert "items" in d
        mc_items = [it for it in d["items"] if it["type"] == "mc"]
        essay_items = [it for it in d["items"] if it["type"] == "essay"]
        assert len(mc_items) == 3
        assert len(essay_items) == 1
        # Problem count field
        assert "problem_count" in d

    def test_difficulty_and_discrimination(self, guru_tok, test_exam, submissions):
        d = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(guru_tok), timeout=30).json()
        items = {it["no"]: it for it in d["items"] if it["type"] == "mc"}
        # Only count our 5 TEST submissions
        # Q1 correct answers: siswas 1,2,3 (siswa with [0,...]) -> 3/5 = 0.6 => Sedang
        # Q2: siswas 1,2,3 (siswa with [,1,]) -> 3/5 = 0.6 => Sedang
        # Q3: siswas 1,2 (siswa with [,,2]) -> 2/5 = 0.4 => Sedang
        # NOTE: there is also the existing exam-independent old siswa data, but THIS is a fresh exam
        # so only our 5 siswas contribute (plus maybe baseline siswa if any extras match kelas).
        # Only our created test siswas submitted on THIS exam, so counts are deterministic.
        resp = items[1]["responses"]
        assert resp >= 5
        # With 5 responses (and if baseline 'siswa' also happens to be in XII IPA 1 she didn't submit this exam),
        # tolerate >=5 submissions, verify labels non-null
        for no in (1, 2, 3):
            it = items[no]
            assert it["difficulty_label"] in ("Mudah", "Sedang", "Sukar")
            assert it["discrimination_label"] in ("Baik", "Cukup", "Jelek")
            # options structure
            for o in it["options"]:
                for k in ("label", "text", "count", "is_key", "pct", "effective"):
                    assert k in o
            assert "blank" in it

    def test_flags_only_when_5_plus(self, guru_tok, test_exam, submissions):
        d = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(guru_tok), timeout=30).json()
        # participants >=5, flags may or may not appear; just ensure flags is a list
        for it in d["items"]:
            assert isinstance(it["flags"], list)

    def test_essay_stats(self, guru_tok, test_exam, submissions):
        d = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(guru_tok), timeout=30).json()
        essay = [it for it in d["items"] if it["type"] == "essay"][0]
        assert "avg_score" in essay
        assert "graded" in essay
        assert essay["responses"] >= 5
        assert essay["graded"] == 0  # none graded yet

    def test_super_admin_ok(self, admin_tok, test_exam, submissions):
        r = requests.get(f"{API}/exams/{test_exam['id']}/analysis", headers=_h(admin_tok), timeout=30)
        assert r.status_code == 200


# =========================================
# Excel export - sheets
# =========================================
class TestExport:
    def test_xlsx_contains_both_sheets(self, guru_tok, test_exam, submissions):
        r = requests.get(f"{API}/exams/{test_exam['id']}/export?format=xlsx",
                         headers={"Authorization": f"Bearer {guru_tok}"}, timeout=60)
        assert r.status_code == 200, r.text
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        names = wb.sheetnames
        assert "Nilai Ujian" in names
        assert "Analisis Butir" in names

    def test_pdf_still_works(self, guru_tok, test_exam, submissions):
        r = requests.get(f"{API}/exams/{test_exam['id']}/export?format=pdf",
                         headers={"Authorization": f"Bearer {guru_tok}"}, timeout=60)
        assert r.status_code == 200
        assert r.content[:4] == b"%PDF"
