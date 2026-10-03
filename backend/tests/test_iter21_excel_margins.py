"""Iteration 21: Excel margin/theme refactor — content starts at B2, A/row1 are spacers.
Covers: pretty_excel (inventory, social-fund, users, class kas) + weekly_attendance_excel.
Also regression: attendance daily export still returns 200 xlsx."""
import os, io, pytest, requests
from openpyxl import load_workbook

BASE = (os.environ.get("REACT_APP_BACKEND_URL") or "https://sekolah-web-14.preview.emergentagent.com").rstrip("/")
API = f"{BASE}/api"

TU = {"email": "tu.demo@sekolahku.id", "password": "Demo12345"}
SUPER = {"email": "boassibarani123@gmail.com", "password": "Boas12345io"}


def login(creds):
    r = requests.post(f"{API}/auth/login", json=creds, timeout=30)
    assert r.status_code == 200, f"login {creds['email']} failed: {r.status_code} {r.text}"
    return r.json()["token"]


def auth(tok):
    return {"Authorization": f"Bearer {tok}"}


@pytest.fixture(scope="module")
def tu_tok():
    return login(TU)


@pytest.fixture(scope="module")
def super_tok():
    return login(SUPER)


def _assert_margin_layout(ws, expected_brand="LAGUBOTI"):
    """Shared assertion: A1 empty, column A values empty, title at B2 contains brand."""
    # Row 1 must be empty (top margin)
    for cell in ws[1]:
        assert cell.value in (None, ""), f"Row 1 should be empty, found {cell.coordinate}={cell.value!r}"
    # Column A must be empty across all rows (left margin)
    for row in ws.iter_rows(min_col=1, max_col=1):
        for cell in row:
            assert cell.value in (None, ""), f"Col A should be empty, found {cell.coordinate}={cell.value!r}"
    # Title at B2
    b2 = ws["B2"].value
    assert b2 is not None and expected_brand in str(b2).upper(), f"B2 should contain '{expected_brand}', got {b2!r}"
    # Column A width is the small margin width
    assert (ws.column_dimensions["A"].width or 0) <= 4, f"Col A should be a thin margin, got width={ws.column_dimensions['A'].width}"


# ---------------- pretty_excel endpoints ----------------
def test_inventory_export_margins(tu_tok):
    r = requests.get(f"{API}/inventory/export", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    ws = load_workbook(io.BytesIO(r.content)).active
    _assert_margin_layout(ws)
    # Header row should be row 5
    row5 = [c.value for c in ws[5] if c.value]
    assert row5, "Row 5 (header) should have content"
    # Summary labels fully visible OR empty-state placeholder
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert ("Total Item" in text) or ("Belum ada data" in text)


def test_social_fund_export_margins(tu_tok):
    r = requests.get(f"{API}/social-fund/export", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    _assert_margin_layout(ws)
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    # If data present, labels; if empty, placeholder
    if "Belum ada data" not in text:
        for lab in ["Pemasukan (Rp)", "Pengeluaran (Rp)", "Saldo Akhir (Rp)", "Total Transaksi"]:
            assert lab in text, f"missing/truncated label '{lab}'"


def test_users_export_margins(super_tok):
    r = requests.get(f"{API}/users/export", headers=auth(super_tok), timeout=30)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    _assert_margin_layout(ws)


def test_class_kas_export_margins(tu_tok):
    # pick any class (TU has visibility of demo classes)
    r = requests.get(f"{API}/classes", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200
    classes = r.json()
    if not classes:
        pytest.skip("no classes available")
    cid = classes[0]["id"]
    r = requests.get(f"{API}/classes/{cid}/kas/export", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    _assert_margin_layout(ws)


# ---------------- weekly_attendance_excel ----------------
def test_weekly_attendance_export_margins_and_alignment(tu_tok):
    r = requests.get(f"{API}/attendance/week/export", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    _assert_margin_layout(ws)
    # Header row at 5. Columns: No, Nama Siswa, Kelas + 7 days + H/I/S/A -> 14 cols, starting at B
    hdr = [c.value for c in ws[5]]
    # Build cleaned mapping col-letter -> text
    assert hdr[0] is None, "A5 should be empty (margin)"
    assert hdr[1] is not None and "No" in str(hdr[1]), f"B5 should be 'No', got {hdr[1]!r}"
    assert "Nama" in str(hdr[2]), f"C5 should contain 'Nama', got {hdr[2]!r}"
    assert "Kelas" in str(hdr[3]), f"D5 should contain 'Kelas', got {hdr[3]!r}"
    # Day headers at E..K (CO=2 + 3..9) = cols 5..11
    for i, dayname in enumerate(["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]):
        v = str(hdr[4 + i] or "")
        assert dayname in v, f"col {4+i} ({chr(ord('A')+4+i)}5) should contain {dayname}, got {v!r}"
    # H/I/S/A totals at next 4 cols
    for i, letter in enumerate(["H", "I", "S", "A"]):
        v = hdr[11 + i]
        assert v == letter, f"total col {11+i} should be {letter}, got {v!r}"
    # Legend keterangan row present
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Keterangan" in text, "legend row should be present"


# ---------------- Regression: daily attendance export ----------------
def test_attendance_daily_export_still_works(tu_tok):
    r = requests.get(f"{API}/attendance/export", headers=auth(tu_tok), timeout=30)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    # Valid xlsx
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row >= 2


# ---------------- PPTX endpoint still works (not changed) ----------------
def test_presentation_pptx_still_200():
    r = requests.get(f"{API}/presentation/pptx", timeout=60)
    assert r.status_code == 200
    assert "presentationml" in r.headers.get("content-type", "") or r.headers.get("content-type", "").startswith("application/")
